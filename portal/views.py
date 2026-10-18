"""
Views for the NIELIT student portal.
Filtering lives in portal/filters.py, Excel import in portal/importer.py.
"""

import json
import re
import traceback
from datetime import datetime
from decimal import Decimal, InvalidOperation

import openpyxl
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.db.models import Count
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect, render
from django.views.decorators.csrf import ensure_csrf_cookie
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .filters import apply_filters
from .forms import ExcelUploadForm, StudentDataForm
from .importer import import_workbook
from .models import Dlc, NsqfElectronics, NsqfIT, studentdata

MONTH_MAP = {
    m: i
    for i, m in enumerate(
        [
            "JAN",
            "FEB",
            "MAR",
            "APR",
            "MAY",
            "JUN",
            "JUL",
            "AUG",
            "SEP",
            "OCT",
            "NOV",
            "DEC",
        ],
        1,
    )
}

SORTABLE_FIELDS = {
    "roll_number",
    "batch_code",
    "name",
    "course_name",
    "father_name",
    "mother_name",
    "dob",
    "gender",
    "address",
    "qualifications",
    "aadhaar",
    "scheme",
    "nsqf",
    "course_hour",
    "course_category",
    "center_name",
    "mode",
    "caste_category",
    "fee",
    "claimable_amount",
    "fee_date",
    "trained",
    "certified",
    "placed",
    "session",
    "claimed",
}

CENTERS = ["inderlok", "janakpuri", "karkardooma"]


def parse_bool(value):
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in ["true", "yes", "1"]


def parse_date(value):
    """Return a date object or None. Accepts openpyxl date/datetime, or common string formats."""
    if not value:
        return None
    # openpyxl may give datetime/date objects
    if hasattr(value, "date"):
        try:
            return value.date()
        except Exception:
            pass
    val_str = str(value).strip()
    # try multiple formats
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(val_str, fmt).date()
        except Exception:
            continue
    # fallback: None
    return None


def format_session_date(value):
    """Return MON-YYYY (e.g. JAN-2024) or empty string."""
    if not value:
        return ""
    # If it's a date/datetime, convert to MON-YYYY
    if hasattr(value, "strftime"):
        return value.strftime("%b-%Y").upper()
    # If it's already a string like JAN-2024 or Jan-2024 or 2024-01, normalize
    s = str(value).strip()
    # If looks like e.g. '2024-01' convert to 'JAN-2024'
    m = re.match(r"^(\d{4})[-/](\d{1,2})$", s)
    if m:
        year = m.group(1)
        month_no = int(m.group(2))
        # reverse lookup month abbrev
        month_abbr = (
            list(MONTH_MAP.keys())[month_no - 1] if 1 <= month_no <= 12 else None
        )
        if month_abbr:
            return f"{month_abbr}-{year}"
    # If already MON-YYYY (or similar), upper and normalize hyphen
    s2 = s.replace(" ", "-").replace("/", "-").upper()
    return s2


def quarter_from_date(date_str):
    if not date_str:
        return None, None
    try:
        month_str, year_str = date_str.upper().split("-")[:2]
        month = MONTH_MAP.get(month_str)
        return (f"Q{(month - 1) // 3 + 1}" if month else None), year_str
    except Exception:
        return None, None


def student_to_dict(s, selected_quarter=None, year=None):
    # Calculate claimable amount based on selected quarter if provided
    if selected_quarter:
        claimable = float(s.get_claimable_amount_for_quarter(selected_quarter, year))
    else:
        claimable = float(s.claimable_amount)
    
    return {
        "dob": s.dob.isoformat() if s.dob else "",
        "fee_date": s.fee_date.isoformat() if s.fee_date else "",
        "trained_date": s.trained_date.isoformat() if s.trained_date else "",
        "session": s.session or "",
        "id": s.id,
        "roll_number": s.roll_number,
        "batch_code": s.batch_code,
        "name": s.name,
        "father_name": s.father_name,
        "mother_name": s.mother_name,
        "gender": s.gender,
        "address": s.address,
        "qualifications": s.qualifications,
        "aadhaar": s.aadhaar,
        "course_name": s.course_name,
        "scheme": s.scheme,
        "nsqf": s.nsqf,
        "course_hour": s.course_hour,
        "course_category": s.course_category,
        "center_name": s.center_name,
        "mode": s.mode,
        "caste_category": s.caste_category,
        "fee": float(s.fee),
        "claimable_amount": claimable,
        "trained": s.trained,
        "certified": s.certified,
        "certified_date": s.certified_date.isoformat() if s.certified_date else "",
        "placed": s.placed,
        "claimed": s.claimed,
    }


def xlrow_to_dict(s):
    return {
        "roll_number": s.roll_number,
        "batch_code": s.batch_code,
        "name": s.name,
        "father_name": s.father_name,
        "mother_name": s.mother_name,
        "dob": s.dob.strftime("%Y-%m-%d") if s.dob else "",
        "gender": s.gender,
        "address": s.address,
        "qualifications": s.qualifications,
        "aadhaar": s.aadhaar,
        "course_name": s.course_name,
        "scheme": s.scheme,
        "nsqf": s.nsqf,
        "course_hour": s.course_hour,
        "center_name": s.center_name,
        "mode": s.mode,
        "caste_category": s.caste_category,
        "fee": float(s.fee),
        "fee_date": s.fee_date or "",
        "trained_date": s.trained_date.isoformat() if s.trained_date else "",
        "certified_date": s.certified_date.isoformat() if s.certified_date else "",
        "placed": s.placed,
        "claimed": s.claimed,
        "session": s.session,
    }


def center_summary(qs):
    summary = {
        "Total": qs.count(),
        "SC": 0, "ST": 0, "OBC": 0, "PWD": 0, "GENERAL": 0,
        "B": 0, "C": 0, "D": 0, "E": 0,
    }
    for row in qs.values("caste_category").annotate(n=Count("id")):
        if row["caste_category"] in summary:
            summary[row["caste_category"]] += row["n"]
    for row in qs.values("course_category").annotate(n=Count("id")):
        cat = (row["course_category"] or "").strip().upper()[:1]
        if cat in summary:
            summary[cat] += row["n"]
    return summary


def login_view(request):
    if request.method == "POST":
        user = authenticate(
            request,
            username=request.POST.get("username", ""),
            password=request.POST.get("password", ""),
        )
        if user:
            login(request, user)
            return redirect("dashboard")
        return render(request, "login.html", {"error": "Invalid credentials"})
    return render(request, "login.html")


def logout_view(request):
    logout(request)
    return redirect("login")


@login_required(login_url="/login")
@ensure_csrf_cookie
def dashboard(request):
    return render(request, "dashboard.html")


@login_required(login_url="/login")
def upload(request):
    """Upload an Excel file. Values are cleaned by portal/importer.py (same code as
    `manage.py import_students`). session comes from the page dropdowns as MON-YYYY."""
    if request.method == "POST":
        form = ExcelUploadForm(request.POST, request.FILES)
        if form.is_valid():
            uploaded_file = request.FILES.get("file")
            month = (request.POST.get("session") or "").strip().upper()
            year = (request.POST.get("year") or "").strip()
            form_session = f"{month}-{year}" if month and year else ""
            try:
                ws = openpyxl.load_workbook(uploaded_file, data_only=True).active
                result = import_workbook(ws, session=form_session)
                messages.success(
                    request,
                    f"Uploaded: {result.created} | Duplicates skipped: {result.duplicates} "
                    f"| Errors: {len(result.bad_rows)}",
                )
                for text, n in result.warnings.items():
                    messages.warning(request, f"{n} rows: {text}")
                for msg in result.bad_rows[:5]:
                    messages.warning(request, msg)
            except Exception as e:
                messages.error(request, f"Error processing file: {e}")
    else:
        form = ExcelUploadForm()

    return render(request, "upload.html", {"form": form})


@login_required(login_url="/login")
def filter_students(request):
    students = apply_filters(request.GET)
    selected_quarter = request.GET.get("quarterly") or None
    year = request.GET.get("year") or None

    def as_int(value, default, maximum=None):
        try:
            n = int(value)
        except (TypeError, ValueError):
            return default
        n = max(n, 1)
        return min(n, maximum) if maximum else n

    page = as_int(request.GET.get("page"), 1)
    limit = as_int(request.GET.get("limit"), 10, maximum=200)
    offset = (page - 1) * limit
    total = students.count()
    total_pages = (total + limit - 1) // limit

    return JsonResponse(
        {
            "results": [
                student_to_dict(s, selected_quarter, year)
                for s in students[offset : offset + limit]
            ],
            "pagination": {
                "page": page,
                "limit": limit,
                "total_count": total,
                "total_pages": total_pages,
                "has_next": page < total_pages,
                "has_prev": page > 1,
            },
        }
    )


@login_required(login_url="/login")
def download_filtered_data(request):
    students = apply_filters(request.GET)
    selected_quarter = request.GET.get("quarterly") or None
    year = request.GET.get("year") or None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Filtered Students"

    headers = [
        "Roll Number",
        "Batch Code",
        "Name",
        "Father Name",
        "Mother Name",
        "DOB",
        "Gender",
        "Address",
        "Qualifications",
        "Aadhaar",
        "Course Name",
        "Scheme",
        "NSQF",
        "Course Hours",
        "Course Category",
        "Center",
        "Mode",
        "Caste Category",
        "Fee",
        "Claimable Amount",
        "Fee Date",
        "Trained",
        "Trained Date",
        "Certified",
        "Certified Date",
        "Placed",
        "Claimed",
        "Session",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            start_color="CCCCCC", end_color="CCCCCC", fill_type="solid"
        )

    yn = lambda v: "Yes" if v else "No"
    for row, s in enumerate(students.iterator(), 2):
        # Calculate claimable amount based on selected quarter if provided
        if selected_quarter:
            claimable = s.get_claimable_amount_for_quarter(selected_quarter, year)
        else:
            claimable = s.claimable_amount
        
        for col, val in enumerate(
            [
                s.roll_number,
                s.batch_code,
                s.name,
                s.father_name,
                s.mother_name,
                s.dob.strftime("%Y-%m-%d") if s.dob else "",
                s.gender,
                s.address,
                s.qualifications,
                s.aadhaar,
                s.course_name,
                s.scheme,
                s.nsqf,
                s.course_hour,
                s.course_category,
                s.center_name,
                s.mode,
                s.caste_category,
                float(s.fee),
                float(claimable),
                str(s.fee_date) if s.fee_date else "",
                yn(s.trained),
                s.trained_date,
                yn(s.certified),
                s.certified_date,
                yn(s.placed),
                yn(s.claimed),
                s.session,
            ],
            1,
        ):
            ws.cell(row=row, column=col, value=val)

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            max((len(str(c.value)) for c in col if c.value), default=0) + 2, 50
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=filtered_students.xlsx"
    wb.save(response)
    return response


# ─── Report ──────────────────────────────────────────────────────────────────


def _session_filter_options():
    sessions = list(
        studentdata.objects.values_list("session", flat=True)
        .distinct()
        .order_by("-session")
    )
    years = sorted({s.split("-")[1] for s in sessions if "-" in s}, reverse=True)
    months = sorted({s.split("-")[0] for s in sessions if "-" in s})
    return years, months


@login_required(login_url="/login")
def download(request):
    p = request.GET
    students = studentdata.objects.all()
    for key in ["year", "session"]:
        if p.get(key):
            students = students.filter(session__icontains=p[key])
    if p.get("center"):
        students = students.filter(center_name=p["center"])

    grouped = {}
    for s in students:
        key = f"{s.course_category}|{s.course_name}|{s.center_name}|{s.session}"
        if key not in grouped:
            grouped[key] = {
                "category": s.course_category,
                "course_name": s.course_name,
                "course_hour": s.course_hour,
                "center_name": s.center_name,
                "session": s.session,
                "scheme": s.scheme,
                "nsqf": s.nsqf,
                **{
                    c: {"trained": 0, "certified": 0, "placed": 0, "total": 0}
                    for c in ["GENERAL", "OBC", "SC", "ST", "PWD"]
                },
            }
        g = grouped[key][s.caste_category]
        g["total"] += 1
        if s.trained_date:
            g["trained"] += 1
        if s.certified_date:
            g["certified"] += 1
        if s.placed:
            g["placed"] += 1

    report_data = list(grouped.values())
    castes = ["GENERAL", "OBC", "SC", "ST", "PWD"]
    totals = {
        c: {"trained": 0, "certified": 0, "placed": 0, "total": 0} for c in castes
    }
    for item in report_data:
        for c in castes:
            for k in totals[c]:
                totals[c][k] += item[c][k]
    totals["grand_total"] = sum(totals[c]["total"] for c in castes)

    years, months = _session_filter_options()
    return render(
        request,
        "download.html",
        {
            "data": report_data,
            "totals": totals,
            "selected_year": p.get("year", ""),
            "selected_session": p.get("session", ""),
            "selected_center": p.get("center", ""),
            "years": years,
            "months": months,
            "centers": CENTERS,
        },
    )


@login_required(login_url="/login")
def api_download_data(request):
    p = request.GET
    students = studentdata.objects.all()
    for key in ["year", "session"]:
        if p.get(key):
            students = students.filter(session__icontains=p[key])
    if p.get("center"):
        students = students.filter(center_name=p["center"])

    return JsonResponse(
        {
            "results": [
                {
                    "course_category": s.course_category,
                    "course_name": s.course_name,
                    "course_hour": s.course_hour,
                    "center_name": s.center_name,
                    "scheme": s.scheme,
                    "nsqf": s.nsqf,
                    "session": s.session,
                    "caste_category": s.caste_category,
                    "trained_date": s.trained_date,
        "certified_date": s.certified_date.isoformat() if s.certified_date else "",
                    "placed": s.placed,
                    "fee": float(s.fee),
                    "claimable_amount": float(s.claimable_amount),
                }
                for s in students
            ]
        }
    )


# ─── Update ──────────────────────────────────────────────────────────────────


@login_required(login_url="/login")
def update_student(request, student_id):
    if request.method != "POST":
        return JsonResponse({"error": "POST required"}, status=405)

    try:
        student = studentdata.objects.get(id=student_id)
    except studentdata.DoesNotExist:
        return JsonResponse({"error": "Student not found"}, status=404)

    try:
        body = json.loads(request.body)
        current_month = (
            datetime.now().strftime("%b").upper() + "-" + datetime.now().strftime("%Y")
        )

        str_fields = [
            "name",
            "father_name",
            "mother_name",
            "address",
            "qualifications",
            "aadhaar",
            "course_name",
            "scheme",
            "nsqf",
        ]
        aadhaar_in = body.get("aadhaar")
        if (
            aadhaar_in is not None
            and str(aadhaar_in).strip()
            and not re.fullmatch(r"\d{12}", str(aadhaar_in).strip())
        ):
            return JsonResponse(
                {"success": False, "error": "Aadhaar must be exactly 12 digits"},
                status=400,
            )

        for f in str_fields:
            if body.get(f) is not None:
                setattr(student, f, str(body[f]).strip())

        session_value = body.get("session")
        if session_value is not None:
            session_value = str(session_value).strip().upper()
            if re.fullmatch(
                r"(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)-\d{4}",
                session_value,
            ):
                student.session = session_value

        if body.get("batch_code"):
            student.batch_code = str(body["batch_code"]).strip().upper()

        for f in ["dob", "fee_date"]:
            if f in body:
                value = body[f]
                parsed = parse_date(value) if value else None
                if value and parsed is None:
                    raise ValueError(f"Invalid date for {f}: {value}")
                setattr(student, f, parsed)

        for f in ["gender", "mode", "caste_category", "center_name"]:
            if body.get(f) is not None:
                setattr(student, f, body[f])

        try:
            student.course_hour = int(
                body.get("course_hour") or student.course_hour or 0
            )
        except (ValueError, TypeError):
            student.course_hour = 0

        if body.get("fee") not in (None, ""):
            try:
                student.fee = Decimal(str(body["fee"]))
            except InvalidOperation:
                raise ValueError(f"Invalid fee: {body['fee']}")

        if body.get("placed") is not None:
            student.placed = body["placed"]

        for field in ["trained", "certified"]:
            new_val = body.get(field)
            if new_val is not None:
                date_field = f"{field}_date"
                # Check if an explicit date was provided in the body
                explicit_date_str = body.get(date_field)
                if explicit_date_str:
                    parsed_date = parse_date(explicit_date_str)
                    setattr(student, date_field, parsed_date)
                elif new_val and not getattr(student, date_field):
                    # If toggling to True and no date exists, set to today
                    setattr(student, date_field, datetime.now().date())
                elif not new_val:
                    # If toggling to False, clear the date
                    setattr(student, date_field, None)
                setattr(student, field, new_val)

        if body.get("claimed") is not None:
            student.claimed = body["claimed"]

        student.save()
        return JsonResponse(
            {
                "success": True,
                "course_category": student.course_category,
                "claimable_amount": float(student.claimable_amount),
                "trained_date": student.trained_date,
                "certified_date": student.certified_date,
                "fee_date": student.fee_date,
                "dob": student.dob,
                "claimed": student.claimed,
            }
        )

    except json.JSONDecodeError as e:
        return JsonResponse(
            {"success": False, "error": f"Invalid JSON: {e}"}, status=400
        )
    except ValueError as e:
        return JsonResponse({"success": False, "error": str(e)}, status=400)
    except Exception as e:
        print(traceback.format_exc())
        return JsonResponse({"success": False, "error": str(e)}, status=400)


@login_required(login_url="/login")
def inputView(request):
    if request.method == "POST":
        form = StudentDataForm(request.POST)
        if form.is_valid():
            student = form.save(commit=False)
            if student.trained and not student.trained_date:
                student.trained_date = datetime.now().date()
            if student.certified and not student.certified_date:
                student.certified_date = datetime.now().date()
            student.save()
            return redirect("dashboard")
    else:
        form = StudentDataForm()

    return render(
        request,
        "input.html",
        {
            "form": form,
            "months": [
                "JAN",
                "FEB",
                "MAR",
                "APR",
                "MAY",
                "JUN",
                "JUL",
                "AUG",
                "SEP",
                "OCT",
                "NOV",
                "DEC",
            ],
            "years": list(range(2020, 2031)),
        },
    )


# ─── Overview ────────────────────────────────────────────────────────────────


def _overview_context(selected_session):
    students = studentdata.objects.all()
    if selected_session:
        students = students.filter(session=selected_session)

    return {
        "all_record": students.count(),
        "centers": [
            {
                "name": n.capitalize(),
                "stats": center_summary(students.filter(center_name=n)),
            }
            for n in CENTERS
        ],
        "sessions": list(
            studentdata.objects.values_list("session", flat=True)
            .distinct()
            .order_by("-session")
        ),
        "selected_session": selected_session,
    }


@login_required(login_url="/login")
def overview(request):
    ctx = _overview_context(request.GET.get("session", ""))
    ctx["centers"] = [(c["name"], c["stats"]) for c in ctx["centers"]]
    return render(request, "overview.html", ctx)


@login_required(login_url="/login")
def overview_data(request):
    return JsonResponse(_overview_context(request.GET.get("session", "")))


@login_required(login_url="/login")
def courses(request):
    return render(
        request,
        "view_courses.html",
        {
            "It": NsqfIT.objects.all(),
            "elctro": NsqfElectronics.objects.all(),
            "dlc": Dlc.objects.all(),
        },
    )


SAMPLE_HEADERS = [
    "roll_number", "batch_code", "name", "father_name", "mother_name", "dob", "gender",
    "address", "qualifications", "aadhaar", "course_name", "scheme", "nsqf", "course_hour",
    "mode", "caste_category", "center_name", "fee", "fee_date", "trained_date",
    "certified_date", "placed",
]


@login_required(login_url="/login")
def sample_upload(request):
    """Excel template with dropdowns. The importer accepts both the dropdown labels
    (e.g. 'Off Campus', 'Inderlok') and the stored values ('offline', 'inderlok')."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Data Template"
    for col, header in enumerate(SAMPLE_HEADERS, 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    ws.freeze_panes = "A2"
    column_index = {h: i for i, h in enumerate(SAMPLE_HEADERS, 1)}

    def add_dropdown(field_name, options):
        options = [o for o in options if o]
        formula = '"' + ",".join(options) + '"'
        if not options or len(formula) > 255:  # Excel limit for inline lists
            return
        letter = get_column_letter(column_index[field_name])
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Please select a value from the dropdown list."
        dv.errorTitle = "Invalid Entry"
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}1048576")

    schemes = sorted({v for v in studentdata.objects.values_list("scheme", flat=True) if v})
    if "NON-NSQF" not in schemes:
        schemes.append("NON-NSQF")

    add_dropdown("center_name", [c[1] for c in studentdata.CENTER_CHOICES])
    add_dropdown("mode", [c[1] for c in studentdata.MODE_CHOICES])
    add_dropdown("caste_category", [c[1] for c in studentdata.CASTE_CHOICES])
    add_dropdown("placed", ["True", "False"])
    add_dropdown("gender", [c[1] for c in studentdata.GENDER])
    add_dropdown("nsqf", [c[1] for c in studentdata.NSQF_LEVEL])
    add_dropdown("scheme", schemes)
    add_dropdown("course_name", [c[1] for c in studentdata.COURSE_CHOICES])
    add_dropdown("qualifications", [c[1] for c in studentdata.HIGHEST_QUALIFICATION])

    sample_row = [
        "NIELIT0001", "Batch-2024-001", "Aarav Sharma", "Rajesh Sharma", "Sunita Sharma",
        "1998-05-12", "Male", "12/4 Rohini Delhi", "Graduation (B.Sc / B.Com / BA / BBA)",
        "123456789012", "Course on Computer Concepts (CCC)", "PMKVY", "Level 4", "120",
        "Off Campus", "OBC", "Inderlok", "5000", "2024-01-10", "2024-01-15", "2024-03-20",
        "True",
    ]
    for col, value in enumerate(sample_row, 1):
        ws.cell(row=2, column=col, value=value)
    for col in range(1, len(SAMPLE_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="student_upload_template.xlsx"'
    wb.save(response)
    return response
