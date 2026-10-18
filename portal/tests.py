import io
import json
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from .filters import apply_filters
from .importer import import_workbook
from .models import studentdata


def make(**kw):
    data = dict(
        name="TEST", course_name="Other Short Term Course", course_hour=120,
        center_name="inderlok", caste_category="GENERAL", mode="offline", fee=Decimal("1000"),
    )
    data.update(kw)
    return studentdata.objects.create(**data)


class ClaimableTests(TestCase):
    def test_non_ao_trained_only_is_70_percent(self):
        s = make(trained=True, trained_date=date(2025, 5, 1))
        self.assertEqual(s.claimable_amount, Decimal("700"))

    def test_non_ao_trained_and_certified_is_100_percent(self):
        s = make(trained=True, trained_date=date(2025, 5, 1), certified=True, certified_date=date(2025, 8, 1))
        self.assertEqual(s.claimable_amount, Decimal("1000"))

    def test_ao_level_needs_certification(self):
        s = make(course_name="A-Level Course", course_hour=600, trained=True, trained_date=date(2025, 5, 1))
        self.assertEqual(s.claimable_amount, 0)
        s.certified, s.certified_date = True, date(2025, 9, 1)
        s.save()
        self.assertEqual(s.claimable_amount, Decimal("1000"))

    def test_quarter_amounts_split_between_quarters(self):
        s = make(trained=True, trained_date=date(2025, 5, 1), certified=True, certified_date=date(2025, 8, 1))
        self.assertEqual(s.get_claimable_amount_for_quarter("Q1"), Decimal("700"))
        self.assertEqual(s.get_claimable_amount_for_quarter("Q2"), Decimal("300"))
        self.assertEqual(s.get_claimable_amount_for_quarter("Q3"), 0)

    def test_trained_and_certified_same_quarter_is_full_amount(self):
        # old code returned only 30% here
        s = make(trained=True, trained_date=date(2025, 4, 10), certified=True, certified_date=date(2025, 6, 10))
        self.assertEqual(s.get_claimable_amount_for_quarter("Q1"), Decimal("1000"))

    def test_quarter_amount_respects_year(self):
        # trained Q1 2024, certified Q1 2025: in 2024 only the 70% counts
        s = make(trained=True, trained_date=date(2024, 5, 1), certified=True, certified_date=date(2025, 5, 1))
        self.assertEqual(s.get_claimable_amount_for_quarter("Q1", 2024), Decimal("700"))
        self.assertEqual(s.get_claimable_amount_for_quarter("Q1", 2025), Decimal("300"))

    def test_aadhaar_must_be_12_digits(self):
        s = make(aadhaar="123")
        with self.assertRaises(ValidationError):
            s.full_clean()


class FilterTests(TestCase):
    def setUp(self):
        # trained Q1-2023, certified Q1-2024  (the old code mixed these up)
        self.cross = make(name="CROSS", trained=True, trained_date=date(2023, 5, 1),
                          certified=True, certified_date=date(2024, 5, 1))
        self.q1_2024 = make(name="ALICE SHARMA", roll_number="R1", trained=True, trained_date=date(2024, 5, 1))
        self.q3_2024 = make(name="BOB", roll_number="R2", trained=True, trained_date=date(2024, 11, 1),
                            certified=True, certified_date=date(2025, 2, 1), nsqf="Level 4", center_name="janakpuri")
        self.none = make(name="NEVER", nsqf="")

    def names(self, **params):
        return {s.name for s in apply_filters(params)}

    def test_year_and_quarter_use_the_same_date(self):
        # CROSS has Q1 in 2023 (trained) and Q1 in 2024 (certified) -> matches both years legitimately
        self.assertEqual(self.names(quarterly="Q1", year="2024"), {"CROSS", "ALICE SHARMA"})
        # but a Q3 2023 search must NOT match CROSS (month of one date + year of another)
        self.assertEqual(self.names(quarterly="Q3", year="2023"), set())

    def test_trained_yes_with_quarter_checks_trained_date(self):
        self.assertEqual(self.names(trained="true", quarterly="Q1", year="2024"), {"ALICE SHARMA"})

    def test_certified_yes_with_quarter_checks_certified_date(self):
        self.assertEqual(self.names(certified="true", quarterly="Q1", year="2024"), {"CROSS"})
        self.assertEqual(self.names(certified="true", quarterly="Q4", year="2025"), {"BOB"})

    def test_quarter_alone_matches_trained_or_certified(self):
        self.assertEqual(self.names(quarterly="Q4"), {"BOB"})

    def test_nsqf_yes_no(self):
        self.assertEqual(self.names(nsqf="yes"), {"BOB"})
        self.assertEqual(self.names(nsqf="no"), {"CROSS", "ALICE SHARMA", "NEVER"})

    def test_search_name_and_roll(self):
        self.assertEqual(self.names(q="sharma"), {"ALICE SHARMA"})
        self.assertEqual(self.names(q="r2"), {"BOB"})
        self.assertEqual(self.names(search="NEVER"), {"NEVER"})

    def test_center_and_boolean_filters(self):
        self.assertEqual(self.names(center="janakpuri"), {"BOB"})
        self.assertEqual(self.names(trained="false"), {"NEVER"})
        self.assertEqual(self.names(trained="or"), {"CROSS", "ALICE SHARMA", "BOB"})

    def test_sort_params_sent_by_dashboard(self):
        asc = [s.name for s in apply_filters({"sort_field": "name", "sort_order": "asc"})]
        desc = [s.name for s in apply_filters({"sort_field": "name", "sort_order": "desc"})]
        self.assertEqual(asc, sorted(asc))
        self.assertEqual(desc, sorted(asc, reverse=True))
        self.assertEqual(len(apply_filters({"sort_field": "id; DROP TABLE"})), 4)  # unknown field ignored


def workbook(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["roll_number", "batch_code", "name", "gender", "dob", "aadhaar", "course_name", "course_hour",
               "scheme", "nsqf", "mode", "caste_category", "center_name", "fee", "trained_date", "certified_date"])
    for r in rows:
        ws.append(r)
    return wb


GOOD = ["1", "B1", "Test One", "F", "2000-01-02", 111122223333.0, "Course on Computer Concepts (CCC)", 80,
        "PMKVY", "4", "Off Campus", "gen", "Inderlok", 5000, "2025-05-01", None]


class ImporterTests(TestCase):
    def test_labels_are_converted_to_stored_values(self):
        r = import_workbook(workbook([GOOD]).active, session="JAN-2025")
        self.assertEqual((r.created, r.duplicates, len(r.bad_rows)), (1, 0, 0))
        s = studentdata.objects.get()
        self.assertEqual((s.center_name, s.mode, s.gender, s.caste_category, s.nsqf),
                         ("inderlok", "offline", "Female", "GENERAL", "Level 4"))
        self.assertEqual(s.aadhaar, "111122223333")
        self.assertEqual(s.session, "JAN-2025")
        self.assertEqual(s.claimable_amount, Decimal("3500"))
        self.assertEqual(s.course_category, "E - DLC")

    def test_duplicates_and_bad_rows(self):
        bad = ["3", "B1", "", "M", None, None, "x", 1, None, None, None, None, "Nowhere", 0, None, None]
        r = import_workbook(workbook([GOOD, GOOD, bad]).active)
        self.assertEqual((r.created, r.duplicates, len(r.bad_rows)), (1, 1, 1))
        r2 = import_workbook(workbook([GOOD]).active)
        self.assertEqual((r2.created, r2.duplicates), (0, 1))

    def test_same_person_two_courses_is_allowed(self):
        other = list(GOOD)
        other[6] = "A-Level Course"
        r = import_workbook(workbook([GOOD, other]).active)
        self.assertEqual(r.created, 2)

    def test_dry_run_saves_nothing(self):
        r = import_workbook(workbook([GOOD]).active, dry_run=True)
        self.assertEqual(r.created, 1)
        self.assertEqual(studentdata.objects.count(), 0)

    def test_no_invented_dates_and_blank_nsqf_is_null(self):
        row = list(GOOD)
        row[9], row[14] = None, None
        row.append(None)
        wb = workbook([row])
        wb.active.cell(row=1, column=17, value="trained")
        wb.active.cell(row=2, column=17, value="yes")
        r = import_workbook(wb.active)
        s = studentdata.objects.get()
        self.assertTrue(s.trained)
        self.assertIsNone(s.trained_date)
        self.assertIsNone(s.nsqf)
        self.assertTrue(any("trained but no trained_date" in w for w in r.warnings))


class ViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("tester", password="pw12345")
        self.a = make(name="ALICE", trained=True, trained_date=date(2024, 5, 1),
                      certified=True, certified_date=date(2024, 8, 1))
        for i in range(12):
            make(name=f"BULK{i}")

    def login(self):
        self.client.force_login(self.user)

    def test_login_required_everywhere(self):
        for url in ["/", "/filter-students/", "/download-filtered-data/", "/upload/", "/sample/",
                    "/overview/", "/view_courses/", "/download/"]:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 302, url)
            self.assertIn("/login", r["Location"], url)

    def test_filter_students_pagination_and_claimable_per_quarter(self):
        self.login()
        data = self.client.get("/filter-students/", {"limit": 5, "page": 1}).json()
        self.assertEqual(data["pagination"]["total_count"], 13)
        self.assertEqual(data["pagination"]["total_pages"], 3)
        self.assertEqual(len(data["results"]), 5)
        q1 = self.client.get("/filter-students/", {"quarterly": "Q1", "year": "2024", "q": "alice"}).json()
        self.assertEqual([x["name"] for x in q1["results"]], ["ALICE"])
        self.assertEqual(q1["results"][0]["claimable_amount"], 700.0)
        q2 = self.client.get("/filter-students/", {"quarterly": "Q2", "year": "2024"}).json()
        self.assertEqual(q2["results"][0]["claimable_amount"], 300.0)

    def test_bad_page_and_limit_do_not_crash(self):
        self.login()
        r = self.client.get("/filter-students/", {"page": "abc", "limit": "-5"})
        self.assertEqual(r.status_code, 200)
        r = self.client.get("/filter-students/", {"limit": "99999"})
        self.assertEqual(r.json()["pagination"]["limit"], 200)

    def test_download_excel_has_search_and_quarter_applied(self):
        self.login()
        r = self.client.get("/download-filtered-data/", {"q": "alice", "quarterly": "Q1", "year": "2024"})
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 2)
        headers = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(row=2, column=headers.index("Claimable Amount") + 1).value, 700)

    def test_sample_template_downloads_and_reimports(self):
        self.login()
        r = self.client.get("/sample/")
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.cell(row=1, column=3).value, "name")
        res = import_workbook(ws)  # its own sample row (labels like "Off Campus") must import cleanly
        self.assertEqual((res.created, len(res.bad_rows)), (1, 0))

    def test_upload_view(self):
        self.login()
        buf = io.BytesIO()
        workbook([GOOD]).save(buf)
        f = SimpleUploadedFile("s.xlsx", buf.getvalue())
        r = self.client.post("/upload/", {"session": "JAN", "year": "2025", "file": f})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(studentdata.objects.filter(name="Test One", session="JAN-2025").exists())

    def test_update_student(self):
        self.login()
        url = f"/update-student/{self.a.id}/"
        post = lambda body: self.client.post(url, json.dumps(body), content_type="application/json")
        r = post({"fee": "2000", "dob": "", "fee_date": "2024-01-02"})
        self.assertEqual(r.status_code, 200)
        self.a.refresh_from_db()
        self.assertEqual(self.a.fee, Decimal("2000"))
        self.assertEqual(self.a.claimable_amount, Decimal("2000"))  # recalculated on save
        self.assertIsNone(self.a.dob)
        self.assertEqual(post({"fee": 0}).status_code, 200)  # zero fee is now allowed
        self.a.refresh_from_db()
        self.assertEqual(self.a.fee, 0)
        self.assertEqual(post({"aadhaar": "12"}).status_code, 400)
        self.assertEqual(post({"dob": "not-a-date"}).status_code, 400)
        self.assertEqual(post({"fee": "abc"}).status_code, 400)
        self.assertEqual(self.client.get(url).status_code, 405)

    def test_overview_counts(self):
        self.login()
        data = self.client.get("/overview-data/").json()
        inderlok = [c for c in data["centers"] if c["name"] == "Inderlok"][0]["stats"]
        self.assertEqual(inderlok["Total"], 13)
        self.assertEqual(inderlok["GENERAL"], 13)
        self.assertEqual(inderlok["C"], 13)

    def test_login_view_missing_fields_does_not_crash(self):
        self.assertEqual(self.client.post("/login/", {}).status_code, 200)
        r = self.client.post("/login/", {"username": "tester", "password": "pw12345"})
        self.assertEqual(r.status_code, 302)
